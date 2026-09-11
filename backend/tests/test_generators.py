from hashlib import sha256
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.generator import _generate_docx, _generate_xlsx, sanitize_output_name
from app.models import MappingItem, MappingMode


def digest(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def test_docx_source_is_immutable_and_detail_row_repeats(tmp_path: Path):
    source = tmp_path / "虚构合同模板.docx"
    destination = tmp_path / "结果.docx"
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("供应商占位")
    run.bold = True
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "商品"
    table.cell(0, 1).text = "数量"
    table.cell(1, 0).text = "商品占位"
    table.cell(1, 1).text = "数量占位"
    document.save(source)
    before = digest(source)

    items = [
        MappingItem(target_id="p:0", source_column="供应商", mode=MappingMode.first_non_empty),
        MappingItem(target_id="t:0:r:1:c:0", source_column="商品", mode=MappingMode.detail),
        MappingItem(target_id="t:0:r:1:c:1", source_column="数量", mode=MappingMode.detail),
    ]
    rows = [
        {"供应商": "虚构医疗用品公司", "商品": "测试手套", "数量": 2},
        {"供应商": "虚构医疗用品公司", "商品": "测试口罩", "数量": 3},
    ]
    _generate_docx(source, destination, items, rows)

    result = Document(destination)
    assert digest(source) == before
    assert result.paragraphs[0].text == "虚构医疗用品公司"
    assert result.paragraphs[0].runs[0].bold
    assert len(result.tables[0].rows) == 3
    assert result.tables[0].cell(1, 0).text == "测试手套"
    assert result.tables[0].cell(2, 1).text == "3"


def test_xlsx_preserves_other_sheet_and_repeats_styled_row(tmp_path: Path):
    source = tmp_path / "虚构合同模板.xlsx"
    destination = tmp_path / "结果.xlsx"
    workbook = Workbook()
    contract = workbook.active
    contract.title = "合同"
    contract["A1"] = "供应商占位"
    contract["A3"] = "商品占位"
    contract["B3"] = "数量占位"
    contract["A3"].font = Font(bold=True, color="FFFFFF")
    contract["A3"].fill = PatternFill("solid", fgColor="1D3D70")
    contract["A3"].alignment = Alignment(horizontal="center")
    contract.row_dimensions[3].height = 25
    contract.print_area = "A1:B4"
    hidden = workbook.create_sheet("内部数据")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "必须保留"
    workbook.save(source)
    before = digest(source)

    items = [
        MappingItem(
            target_id="c:合同:A1", source_column="供应商", mode=MappingMode.first_non_empty
        ),
        MappingItem(target_id="c:合同:A3", source_column="商品", mode=MappingMode.detail),
        MappingItem(target_id="c:合同:B3", source_column="数量", mode=MappingMode.detail),
    ]
    rows = [
        {"供应商": "虚构供应商", "商品": "测试商品A", "数量": 2},
        {"供应商": "虚构供应商", "商品": "测试商品B", "数量": 5},
    ]
    _generate_xlsx(source, destination, "合同", items, rows)

    result = load_workbook(destination, data_only=False)
    assert digest(source) == before
    assert result["合同"]["A1"].value == "虚构供应商"
    assert result["合同"]["A3"].value == "测试商品A"
    assert result["合同"]["A4"].value == "测试商品B"
    assert result["合同"]["A4"].font.bold
    assert result["合同"].row_dimensions[4].height == 25
    assert result["内部数据"]["A1"].value == "必须保留"
    assert result["内部数据"].sheet_state == "hidden"
    result.close()


def test_output_name_is_user_controlled_and_sanitized():
    assert sanitize_output_name("  华东/9月:采购合同  ") == "华东_9月_采购合同"
