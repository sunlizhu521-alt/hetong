import time
import zipfile
from pathlib import Path

import pytest
from docx import Document
from fastapi import HTTPException
from openpyxl import Workbook

from app.file_validation import validate_file
from app.parsers import inspect_order, inspect_template, read_order_rows
from app.storage import SessionStore


def test_xlsx_order_inspection_deduplicates_headers(tmp_path: Path):
    path = tmp_path / "虚构订单.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "订单明细"
    sheet.append(["商品", "商品", "数量", None])
    sheet.append(["测试商品", "规格A", 2, "备注"])
    workbook.save(path)

    inspection = inspect_order(path)
    assert inspection["sheets"][0]["headers"] == ["商品", "商品_2", "数量", "未命名字段4"]
    headers, rows = read_order_rows(path, "订单明细")
    assert headers[1] == "商品_2"
    assert rows[0]["数量"] == 2


def test_docx_template_exposes_paragraph_and_table_targets(tmp_path: Path):
    path = tmp_path / "虚构模板.docx"
    document = Document()
    document.add_paragraph("合同编号")
    document.add_table(rows=1, cols=2)
    document.save(path)

    inspection = inspect_template(path)
    ids = {target["id"] for target in inspection["targets"]}
    assert inspection["kind"] == "docx"
    assert "p:0" in ids
    assert "t:0:r:0:c:0" in ids


def test_expired_session_is_rejected_and_cleaned(tmp_path: Path):
    session_store = SessionStore(tmp_path / "sessions")
    session, token = session_store.create()
    session.meta["expires_at"] = time.time() - 1
    session_store.update_meta(session, expires_at=session.meta["expires_at"])
    with pytest.raises(HTTPException) as error:
        session_store.get(session.id, token)
    assert error.value.status_code == 410
    assert session_store.cleanup_expired() == 1
    assert not session.path.exists()


def test_external_relationship_is_rejected(tmp_path: Path):
    path = tmp_path / "虚构外链模板.xlsx"
    workbook = Workbook()
    workbook.save(path)
    replacement = tmp_path / "rewritten.xlsx"
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(replacement, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "_rels/.rels":
                payload = payload.replace(
                    b"</Relationships>",
                    (
                        b'<Relationship Id="rIdExternal" '
                        b'Type="https://example.invalid/type" '
                        b'Target="https://example.invalid/data" '
                        b'TargetMode="External"/></Relationships>'
                    ),
                )
            target.writestr(info, payload)
    replacement.replace(path)

    with pytest.raises(HTTPException) as error:
        validate_file(path, ".xlsx")
    assert error.value.status_code == 415
    assert "外部链接" in error.value.detail
