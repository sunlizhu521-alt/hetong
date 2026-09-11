from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

import app.main as main_module
from app.storage import SessionStore


def _xlsx_bytes() -> bytes:
    stream = BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "虚构订单"
    sheet.append(["供应商", "商品", "数量"])
    sheet.append(["示例供应商", "示例商品A", 2])
    sheet.append(["示例供应商", "示例商品B", 3])
    workbook.save(stream)
    return stream.getvalue()


def _docx_bytes() -> bytes:
    stream = BytesIO()
    document = Document()
    document.add_heading("虚构采购合同", level=1)
    document.add_paragraph("供应商")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "商品"
    table.cell(0, 1).text = "数量"
    document.save(stream)
    return stream.getvalue()


def test_session_upload_inspect_and_token_isolation(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(main_module, "store", SessionStore(tmp_path / "sessions"))
    monkeypatch.setattr(main_module, "rate_limiter", main_module.CreateRateLimiter())

    with TestClient(main_module.app) as client:
        created = client.post("/hetong-api/v1/sessions")
        assert created.status_code == 200
        session = created.json()
        headers = {"X-Session-Token": session["token"]}

        order = client.request(
            "PUT",
            f"/hetong-api/v1/sessions/{session['id']}/order",
            headers=headers,
            files={
                "file": (
                    "虚构订单.xlsx",
                    _xlsx_bytes(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert order.status_code == 200
        assert order.json()["filename"] == "虚构订单.xlsx"

        template = client.request(
            "PUT",
            f"/hetong-api/v1/sessions/{session['id']}/template",
            headers=headers,
            files={
                "file": (
                    "虚构模板.docx",
                    _docx_bytes(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        assert template.status_code == 200

        inspected = client.post(
            f"/hetong-api/v1/sessions/{session['id']}/inspect",
            headers=headers,
            json={"order_sheet": "虚构订单", "template_sheet": None},
        )
        assert inspected.status_code == 200
        body = inspected.json()
        assert body["order"]["sheets"][0]["name"] == "虚构订单"
        assert body["template"]["kind"] == "docx"
        assert body["template"]["fingerprint"] == template.json()["sha256"]

        denied = client.post(
            f"/hetong-api/v1/sessions/{session['id']}/inspect",
            headers={"X-Session-Token": "wrong-token"},
            json={"order_sheet": "虚构订单", "template_sheet": None},
        )
        assert denied.status_code == 403
