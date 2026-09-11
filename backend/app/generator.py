from __future__ import annotations

import copy
import re
import shutil
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter, range_boundaries

from .mapping import resolve_scalar
from .models import GenerateRequest, MappingItem, MappingMode
from .office import convert_to_pdf, convert_xls_to_xlsx, render_pdf_pages
from .parsers import display_value

INVALID_FILENAME = re.compile(r"[\\/:*?\"<>|\x00-\x1f]")


def sanitize_output_name(name: str) -> str:
    cleaned = INVALID_FILENAME.sub("_", name).strip().strip(".")
    cleaned = re.sub(r"\s+", " ", cleaned)
    if not cleaned:
        raise ValueError("合同名称不能为空")
    return cleaned[:100]


def generate_contract(
    template_path: Path,
    request: GenerateRequest,
    rows: list[dict[str, Any]],
    result_dir: Path,
    preview_dir: Path,
) -> dict:
    result_dir.mkdir(parents=True, exist_ok=True)
    preview_dir.mkdir(parents=True, exist_ok=True)
    output_name = sanitize_output_name(request.output_name)
    source = template_path
    if source.suffix.lower() == ".xls":
        converted_dir = template_path.parent.parent / "work" / "converted"
        source = convert_xls_to_xlsx(source, converted_dir)

    if source.suffix.lower() == ".docx":
        editable = result_dir / f"{output_name}.docx"
        _generate_docx(source, editable, request.items, rows)
        editable_kind = "docx"
    else:
        editable = result_dir / f"{output_name}.xlsx"
        _generate_xlsx(source, editable, request.template_sheet, request.items, rows)
        editable_kind = "xlsx"

    pdf = convert_to_pdf(editable, result_dir)
    desired_pdf = result_dir / f"{output_name}.pdf"
    if pdf != desired_pdf:
        pdf.replace(desired_pdf)
    pages = render_pdf_pages(desired_pdf, preview_dir)
    return {
        "output_name": output_name,
        "editable_kind": editable_kind,
        "files": {editable_kind: editable.name, "pdf": desired_pdf.name},
        "preview_pages": [page.name for page in pages],
    }


def _generate_docx(
    source: Path,
    destination: Path,
    items: list[MappingItem],
    rows: list[dict[str, Any]],
) -> None:
    shutil.copy2(source, destination)
    document = Document(destination)
    scalar_items = [item for item in items if item.mode != MappingMode.detail]
    detail_items = [item for item in items if item.mode == MappingMode.detail]
    for item in scalar_items:
        _set_docx_target(document, item.target_id, resolve_scalar(item, rows))

    groups: dict[tuple[int, int], list[MappingItem]] = {}
    for item in detail_items:
        match = re.fullmatch(r"t:(\d+):r:(\d+):c:(\d+)", item.target_id)
        if not match:
            raise ValueError("Word 明细字段必须映射到表格单元格")
        table_index, row_index, _ = map(int, match.groups())
        groups.setdefault((table_index, row_index), []).append(item)
    for (table_index, row_index), group in sorted(groups.items(), reverse=True):
        table = document.tables[table_index]
        template_row = table.rows[row_index]
        template_xml = copy.deepcopy(template_row._tr)
        anchor_xml = template_row._tr
        for offset, data_row in enumerate(rows):
            if offset:
                new_row_xml = copy.deepcopy(template_xml)
                anchor_xml.addnext(new_row_xml)
                anchor_xml = new_row_xml
            current_row = table.rows[row_index + offset]
            for item in group:
                column_index = int(item.target_id.rsplit(":", 1)[1])
                value = display_value(data_row.get(item.source_column or ""))
                _set_cell_text(current_row.cells[column_index], value)
    document.save(destination)


def _set_docx_target(document: Document, target_id: str, value: str) -> None:
    paragraph_match = re.fullmatch(r"p:(\d+)", target_id)
    if paragraph_match:
        paragraph = document.paragraphs[int(paragraph_match.group(1))]
        _set_paragraph_text(paragraph, value)
        return
    cell_match = re.fullmatch(r"t:(\d+):r:(\d+):c:(\d+)", target_id)
    if not cell_match:
        raise ValueError(f"Word 目标位置无效：{target_id}")
    table_index, row_index, column_index = map(int, cell_match.groups())
    _set_cell_text(document.tables[table_index].rows[row_index].cells[column_index], value)


def _set_cell_text(cell, value: str) -> None:
    paragraph = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
    _set_paragraph_text(paragraph, value)
    for extra in cell.paragraphs[1:]:
        _set_paragraph_text(extra, "")


def _set_paragraph_text(paragraph, value: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)


def _generate_xlsx(
    source: Path,
    destination: Path,
    sheet_name: str | None,
    items: list[MappingItem],
    rows: list[dict[str, Any]],
) -> None:
    shutil.copy2(source, destination)
    workbook = load_workbook(destination, data_only=False, keep_links=False)
    if not sheet_name or sheet_name not in workbook.sheetnames:
        workbook.close()
        raise ValueError("合同模板 Sheet 不存在")
    sheet = workbook[sheet_name]
    scalar_items = [item for item in items if item.mode != MappingMode.detail]
    detail_items = [item for item in items if item.mode == MappingMode.detail]
    for item in scalar_items:
        target_sheet, coordinate = _parse_excel_target(item.target_id)
        if target_sheet != sheet_name:
            workbook.close()
            raise ValueError("映射目标不属于所选合同 Sheet")
        sheet[coordinate] = resolve_scalar(item, rows)

    row_groups: dict[int, list[MappingItem]] = {}
    for item in detail_items:
        target_sheet, coordinate = _parse_excel_target(item.target_id)
        if target_sheet != sheet_name:
            workbook.close()
            raise ValueError("映射目标不属于所选合同 Sheet")
        row_groups.setdefault(sheet[coordinate].row, []).append(item)
    for row_index, group in sorted(row_groups.items(), reverse=True):
        _repeat_excel_row(sheet, row_index, group, rows)
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(destination)
    workbook.close()


def _parse_excel_target(target_id: str) -> tuple[str, str]:
    match = re.fullmatch(r"c:(.+):([A-Z]+\d+)", target_id)
    if not match:
        raise ValueError(f"Excel 目标位置无效：{target_id}")
    return match.group(1), match.group(2)


def _repeat_excel_row(
    sheet, row_index: int, items: list[MappingItem], rows: list[dict[str, Any]]
) -> None:
    if not rows:
        rows = [{}]
    amount = len(rows) - 1
    merged_ranges = list(sheet.merged_cells.ranges)
    for merged in merged_ranges:
        if merged.min_row < row_index < merged.max_row:
            raise ValueError("明细模板行穿过纵向合并单元格，第一版无法安全复制")
    source_snapshot = []
    for column in range(1, sheet.max_column + 1):
        cell = sheet.cell(row=row_index, column=column)
        source_snapshot.append(
            {
                "column": column,
                "value": cell.value,
                "style": copy.copy(cell._style),
                "number_format": cell.number_format,
                "font": copy.copy(cell.font),
                "fill": copy.copy(cell.fill),
                "border": copy.copy(cell.border),
                "alignment": copy.copy(cell.alignment),
                "protection": copy.copy(cell.protection),
            }
        )
    row_height = sheet.row_dimensions[row_index].height
    if amount:
        for merged in merged_ranges:
            sheet.unmerge_cells(str(merged))
        sheet.insert_rows(row_index + 1, amount=amount)
        for merged in merged_ranges:
            min_col, min_row, max_col, max_row = range_boundaries(str(merged))
            if min_row > row_index:
                min_row += amount
                max_row += amount
            if min_row == max_row == row_index:
                for offset in range(len(rows)):
                    sheet.merge_cells(
                        start_row=row_index + offset,
                        start_column=min_col,
                        end_row=row_index + offset,
                        end_column=max_col,
                    )
            else:
                sheet.merge_cells(
                    start_row=min_row,
                    start_column=min_col,
                    end_row=max_row,
                    end_column=max_col,
                )
    for offset, data_row in enumerate(rows):
        target_row = row_index + offset
        if offset:
            sheet.row_dimensions[target_row].height = row_height
            for snapshot in source_snapshot:
                cell = sheet.cell(row=target_row, column=snapshot["column"])
                if isinstance(cell, MergedCell):
                    continue
                value = snapshot["value"]
                if isinstance(value, str) and value.startswith("="):
                    origin = f"{get_column_letter(snapshot['column'])}{row_index}"
                    destination = f"{get_column_letter(snapshot['column'])}{target_row}"
                    value = Translator(value, origin=origin).translate_formula(destination)
                cell.value = value
                cell._style = copy.copy(snapshot["style"])
                cell.font = copy.copy(snapshot["font"])
                cell.fill = copy.copy(snapshot["fill"])
                cell.border = copy.copy(snapshot["border"])
                cell.alignment = copy.copy(snapshot["alignment"])
                cell.protection = copy.copy(snapshot["protection"])
                cell.number_format = snapshot["number_format"]
        for item in items:
            _, coordinate = _parse_excel_target(item.target_id)
            column = sheet[coordinate].column
            sheet.cell(row=target_row, column=column).value = display_value(
                data_row.get(item.source_column or "")
            )
    _extend_print_area(sheet, row_index, amount)


def _extend_print_area(sheet, row_index: int, amount: int) -> None:
    if not amount or not sheet.print_area:
        return
    updated: list[str] = []
    for area in str(sheet.print_area).replace("'", "").split(","):
        if "!" in area:
            _, area = area.split("!", 1)
        min_col, min_row, max_col, max_row = range_boundaries(area.replace("$", ""))
        if max_row >= row_index:
            max_row += amount
        updated.append(
            f"${get_column_letter(min_col)}${min_row}:${get_column_letter(max_col)}${max_row}"
        )
    sheet.print_area = ",".join(updated)
