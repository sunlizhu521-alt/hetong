from __future__ import annotations

import hashlib
import json
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

import xlrd
from docx import Document
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .config import settings


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    return value


def display_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        while chunk := handle.read(1024 * 1024):
            digest.update(chunk)
    return digest.hexdigest()


def inspect_order(path: Path) -> dict:
    if path.suffix.lower() == ".xls":
        book = xlrd.open_workbook(path, on_demand=True)
        sheets = []
        for sheet in book.sheets():
            header_index = _xls_header_index(sheet)
            headers = _dedupe_headers(
                [display_value(sheet.cell_value(header_index, c)) for c in range(sheet.ncols)]
            )
            preview = []
            for r in range(header_index + 1, min(sheet.nrows, header_index + 6)):
                preview.append(
                    {headers[c]: display_value(sheet.cell_value(r, c)) for c in range(sheet.ncols)}
                )
            sheets.append(
                {
                    "name": sheet.name,
                    "headers": headers,
                    "row_count": max(0, sheet.nrows - header_index - 1),
                    "header_row": header_index + 1,
                    "preview": preview,
                }
            )
        return {"sheets": sheets}

    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    sheets = []
    for sheet in workbook.worksheets:
        header_index, raw_headers = _xlsx_header(sheet)
        headers = _dedupe_headers([display_value(value) for value in raw_headers])
        preview = []
        for row in sheet.iter_rows(
            min_row=header_index + 2,
            max_row=min(sheet.max_row, header_index + 6),
            values_only=True,
        ):
            preview.append(
                {
                    headers[i]: json_value(row[i]) if i < len(row) else None
                    for i in range(len(headers))
                }
            )
        sheets.append(
            {
                "name": sheet.title,
                "headers": headers,
                "row_count": max(0, sheet.max_row - header_index - 1),
                "header_row": header_index + 1,
                "preview": preview,
            }
        )
    workbook.close()
    return {"sheets": sheets}


def read_order_rows(path: Path, sheet_name: str) -> tuple[list[str], list[dict[str, Any]]]:
    if path.suffix.lower() == ".xls":
        book = xlrd.open_workbook(path, on_demand=True)
        if sheet_name not in book.sheet_names():
            raise ValueError("订单 Sheet 不存在")
        sheet = book.sheet_by_name(sheet_name)
        header_index = _xls_header_index(sheet)
        headers = _dedupe_headers(
            [display_value(sheet.cell_value(header_index, c)) for c in range(sheet.ncols)]
        )
        if sheet.nrows - header_index - 1 > settings.max_source_rows:
            raise ValueError(f"订单数据超过 {settings.max_source_rows} 行")
        rows = []
        for r in range(header_index + 1, sheet.nrows):
            values = [sheet.cell_value(r, c) for c in range(sheet.ncols)]
            if not any(display_value(value).strip() for value in values):
                continue
            rows.append({headers[c]: json_value(values[c]) for c in range(len(headers))})
        return headers, rows

    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    if sheet_name not in workbook.sheetnames:
        workbook.close()
        raise ValueError("订单 Sheet 不存在")
    sheet = workbook[sheet_name]
    header_index, raw_headers = _xlsx_header(sheet)
    headers = _dedupe_headers([display_value(value) for value in raw_headers])
    if sheet.max_row - header_index - 1 > settings.max_source_rows:
        workbook.close()
        raise ValueError(f"订单数据超过 {settings.max_source_rows} 行")
    rows = []
    for values in sheet.iter_rows(min_row=header_index + 2, values_only=True):
        if not any(display_value(value).strip() for value in values):
            continue
        rows.append(
            {
                headers[i]: json_value(values[i]) if i < len(values) else None
                for i in range(len(headers))
            }
        )
    workbook.close()
    return headers, rows


def inspect_template(path: Path, selected_sheet: str | None = None) -> dict:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _inspect_docx(path)
    if suffix == ".xls":
        return {
            "kind": "xls",
            "fingerprint": sha256_file(path),
            "sheets": [],
            "targets": [],
            "warnings": ["老式 XLS 模板将在服务器转换为 XLSX 后进行映射。"],
            "blocking_errors": [],
            "requires_conversion": True,
        }
    return _inspect_xlsx(path, selected_sheet)


def _inspect_docx(path: Path) -> dict:
    document = Document(path)
    targets: list[dict] = []
    blocks: list[dict] = []
    fonts: set[str] = set()
    for index, paragraph in enumerate(document.paragraphs):
        text = paragraph.text
        for run in paragraph.runs:
            if run.font.name:
                fonts.add(run.font.name)
        if text.strip():
            target = {
                "id": f"p:{index}",
                "label": f"正文 {index + 1}",
                "text": text,
                "type": "paragraph",
            }
            targets.append(target)
            blocks.append(target)
    for table_index, table in enumerate(document.tables):
        rows: list[list[dict]] = []
        for row_index, row in enumerate(table.rows):
            cells: list[dict] = []
            seen: set[int] = set()
            for column_index, cell in enumerate(row.cells):
                identity = id(cell._tc)
                if identity in seen:
                    continue
                seen.add(identity)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        if run.font.name:
                            fonts.add(run.font.name)
                target = {
                    "id": f"t:{table_index}:r:{row_index}:c:{column_index}",
                    "label": (
                        f"表格 {table_index + 1} · "
                        f"第 {row_index + 1} 行第 {column_index + 1} 列"
                    ),
                    "text": cell.text,
                    "type": "table_cell",
                    "table": table_index,
                    "row": row_index,
                    "column": column_index,
                }
                targets.append(target)
                cells.append(target)
            rows.append(cells)
        blocks.append({"type": "table", "table": table_index, "rows": rows})

    warnings: list[str] = []
    if any(
        section.header.paragraphs and any(p.text.strip() for p in section.header.paragraphs)
        for section in document.sections
    ):
        warnings.append("检测到页眉文字，第一版不支持映射页眉字段。")
    if any(
        section.footer.paragraphs and any(p.text.strip() for p in section.footer.paragraphs)
        for section in document.sections
    ):
        warnings.append("检测到页脚文字，第一版不支持映射页脚字段。")
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
        if b"txbxContent" in xml:
            warnings.append("检测到文本框或艺术字，其中的字段不能映射。")
        if any(name.startswith("word/embeddings/") for name in archive.namelist()):
            warnings.append("检测到嵌入对象，生成前必须确认版式预览。")
    return {
        "kind": "docx",
        "fingerprint": sha256_file(path),
        "sheets": [],
        "targets": targets,
        "blocks": blocks,
        "warnings": warnings,
        "blocking_errors": [],
        "fonts": sorted(fonts),
        "requires_conversion": False,
    }


def _inspect_xlsx(path: Path, selected_sheet: str | None) -> dict:
    workbook = load_workbook(path, data_only=False, keep_links=True, read_only=False)
    warnings: list[str] = []
    blocking: list[str] = []
    if getattr(workbook, "_external_links", None):
        blocking.append("模板包含外部链接，第一版为避免外部数据访问而停止处理。")
    if workbook.security.lockStructure or workbook.security.lockWindows:
        blocking.append("工作簿结构受到保护，请先在 Excel 中解除保护。")
    sheet_names = workbook.sheetnames
    sheet_name = (
        selected_sheet
        if selected_sheet in sheet_names
        else (sheet_names[0] if sheet_names else None)
    )
    targets: list[dict] = []
    grid: list[list[dict]] = []
    fonts: set[str] = set()
    if sheet_name:
        sheet = workbook[sheet_name]
        merged_non_topleft = {
            (row, col)
            for merged in sheet.merged_cells.ranges
            for row in range(merged.min_row, merged.max_row + 1)
            for col in range(merged.min_col, merged.max_col + 1)
            if (row, col) != (merged.min_row, merged.min_col)
        }
        max_row = min(max(sheet.max_row, 20), 200)
        max_col = min(max(sheet.max_column, 8), 30)
        for row in range(1, max_row + 1):
            grid_row: list[dict] = []
            for column in range(1, max_col + 1):
                if (row, column) in merged_non_topleft:
                    continue
                cell = sheet.cell(row=row, column=column)
                if cell.font and cell.font.name:
                    fonts.add(cell.font.name)
                target = {
                    "id": f"c:{sheet_name}:{cell.coordinate}",
                    "label": f"{sheet_name} · {cell.coordinate}",
                    "text": display_value(cell.value),
                    "type": "cell",
                    "row": row,
                    "column": column,
                    "coordinate": cell.coordinate,
                    "column_letter": get_column_letter(column),
                }
                targets.append(target)
                grid_row.append(target)
            grid.append(grid_row)
        if sheet.protection.sheet:
            blocking.append(f"Sheet“{sheet_name}”受到保护，请先解除保护。")
    workbook.close()
    return {
        "kind": "xlsx",
        "fingerprint": sha256_file(path),
        "sheets": sheet_names,
        "selected_sheet": sheet_name,
        "targets": targets,
        "grid": grid,
        "warnings": warnings,
        "blocking_errors": blocking,
        "fonts": sorted(fonts),
        "requires_conversion": False,
    }


def _xlsx_header(sheet) -> tuple[int, list[Any]]:
    for row_index, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=min(20, sheet.max_row), values_only=True)
    ):
        if any(display_value(value).strip() for value in row):
            return row_index, list(row)
    return 0, ["字段1"]


def _xls_header_index(sheet) -> int:
    for row in range(min(20, sheet.nrows)):
        if any(display_value(sheet.cell_value(row, col)).strip() for col in range(sheet.ncols)):
            return row
    return 0


def _dedupe_headers(headers: list[str]) -> list[str]:
    result: list[str] = []
    counts: dict[str, int] = {}
    for index, raw in enumerate(headers):
        base = raw.strip() or f"未命名字段{index + 1}"
        counts[base] = counts.get(base, 0) + 1
        result.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return result


def compact_fingerprint_payload(order: dict, template: dict) -> str:
    payload = {"order": order.get("sheets", []), "template": template.get("fingerprint")}
    return hashlib.sha256(
        json.dumps(payload, ensure_ascii=False, sort_keys=True).encode()
    ).hexdigest()
